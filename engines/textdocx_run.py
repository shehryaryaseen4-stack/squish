"""Build an editable DOCX from plain text (used for scanned PDFs after OCR).
Usage: textdocx_run.py output.docx < text   (pages separated by form feed \\f)"""
import sys, re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
pages = sys.stdin.read().split("\f")
for i, page in enumerate(pages):
    # paragraphs are separated by blank lines; lines inside a paragraph are joined
    for para in re.split(r"\n\s*\n", page.strip()):
        text = " ".join(l.strip() for l in para.splitlines() if l.strip())
        if text:
            doc.add_paragraph(text)
    if i < len(pages) - 1:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
doc.save(sys.argv[1])
